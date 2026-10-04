"""Deliverables: Word, Excel and PowerPoint files written to workspace/outputs.

The model passes content as simple markdown-style text, which is easier for a small
model to produce reliably than nested structures:
    ## Heading      - bullet      1. numbered      | a | b | table rows      **bold**
"""
import csv
import io
import re
from datetime import datetime

from ..config import OUTPUTS
from .files import safe_name

_BOLD = re.compile(r"(\*\*.+?\*\*)")
FOOTER = "Generated on-premise by Sovereign AI Workbench. No data left this machine."


def _text(value):
    """Small models sometimes send a list or dict where a string was asked for."""
    if isinstance(value, str):
        return value
    if isinstance(value, list):
        return "\n".join(_text(v) for v in value)
    if isinstance(value, dict):
        return "\n".join(f"## {k}\n{_text(v)}" for k, v in value.items())
    return "" if value is None else str(value)


as_text = _text


def _table_rows(lines):
    rows = []
    for line in lines:
        cells = [c.strip() for c in line.strip().strip("|").split("|")]
        if all(re.fullmatch(r":?-{2,}:?", c) for c in cells):  # markdown separator row
            continue
        rows.append(cells)
    return rows


def _add_runs(paragraph, text):
    for part in _BOLD.split(text):
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            paragraph.add_run(part[2:-2]).bold = True
        elif part:
            paragraph.add_run(part)


def create_word(args, ctx):
    import docx
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, RGBColor

    name = safe_name(args.get("filename"), ".docx")
    title = _text(args.get("title")) or "Document"
    content = _text(args.get("content"))
    if len(content.strip()) < 20:
        return {"ok": False, "output": "The content is empty. Call create_word again with the full text."}

    doc = docx.Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)

    tag = doc.add_paragraph()
    run = tag.add_run("CONFIDENTIAL - INTERNAL USE ONLY")
    run.bold = True
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0xB0, 0x1E, 0x1E)
    doc.add_heading(title, level=0)
    meta = doc.add_paragraph()
    meta.add_run(f"Date: {datetime.now():%d %B %Y}    Prepared with: Sovereign AI Workbench").italic = True

    lines = content.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        stripped = line.strip()
        if not stripped:
            i += 1
        elif stripped.startswith("|"):
            block = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                block.append(lines[i])
                i += 1
            rows = _table_rows(block)
            if rows:
                cols = max(len(r) for r in rows)
                table = doc.add_table(rows=len(rows), cols=cols)
                table.style = "Light Grid Accent 1"
                for r, row in enumerate(rows):
                    for c, cell in enumerate(row):
                        table.cell(r, c).text = cell.replace("**", "")
                        if r == 0:
                            for p in table.cell(r, c).paragraphs:
                                for rn in p.runs:
                                    rn.bold = True
                doc.add_paragraph()
        else:
            m = re.match(r"^(#{1,4})\s+(.*)", stripped)
            if m:
                doc.add_heading(m.group(2).replace("**", ""), level=min(len(m.group(1)), 3))
            elif re.match(r"^[-*•]\s+", stripped):
                _add_runs(doc.add_paragraph(style="List Bullet"), re.sub(r"^[-*•]\s+", "", stripped))
            elif re.match(r"^\d+[.)]\s+", stripped):
                _add_runs(doc.add_paragraph(style="List Number"), re.sub(r"^\d+[.)]\s+", "", stripped))
            else:
                _add_runs(doc.add_paragraph(), stripped)
            i += 1

    foot = doc.add_paragraph()
    foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = foot.add_run(FOOTER)
    fr.italic = True
    fr.font.size = Pt(8)
    fr.font.color.rgb = RGBColor(0x6B, 0x72, 0x80)

    doc.save(str(OUTPUTS / name))
    return {"ok": True, "output": f"Word file created: {name}", "file": name}


def create_excel(args, ctx):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    name = safe_name(args.get("filename"), ".xlsx")
    content = args.get("content")
    if isinstance(content, list) and content and isinstance(content[0], (list, tuple)):
        rows = [list(r) for r in content]
    else:
        text = _text(content).strip()
        if text.startswith("|"):
            rows = _table_rows([ln for ln in text.splitlines() if ln.strip()])
        else:
            rows = [r for r in csv.reader(io.StringIO(text)) if r]
    if not rows:
        return {"ok": False, "output": "No rows were provided. Send CSV text with a header row."}

    def cell_value(v):
        v = v.strip() if isinstance(v, str) else v
        if isinstance(v, str) and not v.startswith("="):
            try:
                return float(v) if "." in v else int(v)
            except ValueError:
                return v
        return v

    wb = Workbook()
    ws = wb.active
    ws.title = (_text(args.get("sheet_name")) or "Sheet1")[:31]
    for r in rows:
        ws.append([cell_value(v) for v in r])
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1D4ED8")
        cell.alignment = Alignment(horizontal="center")
    for idx in range(1, ws.max_column + 1):
        width = max(len(str(c.value)) if c.value is not None else 0 for c in ws[get_column_letter(idx)])
        ws.column_dimensions[get_column_letter(idx)].width = min(max(10, width + 2), 60)
    ws.freeze_panes = "A2"
    wb.save(str(OUTPUTS / name))
    return {"ok": True, "output": f"Excel file created: {name} ({len(rows)} rows)", "file": name}


def create_presentation(args, ctx):
    from pptx import Presentation
    from pptx.util import Pt

    name = safe_name(args.get("filename"), ".pptx")
    title = _text(args.get("title")) or "Presentation"
    content = _text(args.get("content"))

    slides, current = [], None
    for line in content.splitlines():
        s = line.strip()
        if not s:
            continue
        # a slide starts at a '## Title' line, or at a 'Slide 3: Title' line however it is marked up
        m = re.match(r"^#{1,4}\s+(.*)", s) or re.match(r"^\**\s*(slide\s*\d+\s*[:.\-]\s*.*)", s, re.I)
        if m:
            heading = re.sub(r"^slide\s*\d+\s*[:.\-]\s*", "", m.group(1).replace("**", "").strip(), flags=re.I)
            current = {"title": heading or "Slide", "bullets": []}
            slides.append(current)
        else:
            if current is None:
                current = {"title": "Overview", "bullets": []}
                slides.append(current)
            current["bullets"].append(re.sub(r"^([-*•]|\d+[.)])\s+", "", s).replace("**", ""))
    if not slides:
        return {"ok": False, "output": "No slides were provided. Use '## Slide title' then '- bullet' lines."}

    prs = Presentation()
    first = prs.slides.add_slide(prs.slide_layouts[0])
    first.shapes.title.text = title
    first.placeholders[1].text = f"{datetime.now():%d %B %Y}\nConfidential - prepared on-premise"
    for s in slides:
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = s["title"]
        frame = slide.placeholders[1].text_frame
        frame.clear()
        for n, bullet in enumerate(s["bullets"][:8]):
            p = frame.paragraphs[0] if n == 0 else frame.add_paragraph()
            p.text = bullet
            p.font.size = Pt(20)
    prs.save(str(OUTPUTS / name))
    return {"ok": True, "output": f"PowerPoint file created: {name} ({len(slides) + 1} slides)", "file": name}
