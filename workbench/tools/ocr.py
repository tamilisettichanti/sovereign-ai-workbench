"""Text extraction. Scanned pages and images go through on-device Tesseract OCR."""
import io
import re
from pathlib import Path

import pytesseract
from PIL import Image

from ..config import IMAGE_EXTS, TESSERACT_CMD
from .files import resolve_input

pytesseract.pytesseract.tesseract_cmd = TESSERACT_CMD


_TAG = re.compile(r"\b[A-Z]{1,4}-\d{2,4}[A-Z]?\b")


def _prepare(img):
    """Greyscale, and enlarge small images: Tesseract misreads small label text."""
    img = img.convert("L")
    if max(img.size) < 1800:
        img = img.resize((img.width * 2, img.height * 2), Image.LANCZOS)
    return img


def _ocr_image(img, psm=6):
    return pytesseract.image_to_string(_prepare(img), config=f"--psm {psm}")


def extract_tags(path):
    """Equipment and instrument tags (T-101, FCV-101 ...) found on a drawing.

    Two OCR passes are combined: block mode reads captions, sparse mode reads isolated labels.
    """
    with Image.open(path) as img:
        text = _ocr_image(img, 6) + "\n" + _ocr_image(img, 11)
    return sorted(set(_TAG.findall(text)))


def _text_lines(img):
    """Text lines on a drawing with their positions (sparse-text OCR)."""
    data = pytesseract.image_to_data(_prepare(img), config="--psm 11", output_type=pytesseract.Output.DICT)
    lines = {}
    for i, word in enumerate(data["text"]):
        word = word.strip()
        if not word or float(data["conf"][i]) < 30:
            continue
        key = (data["block_num"][i], data["par_num"][i], data["line_num"][i])
        x0, y0 = data["left"][i], data["top"][i]
        x1, y1 = x0 + data["width"][i], y0 + data["height"][i]
        ln = lines.setdefault(key, {"words": [], "x0": x0, "y0": y0, "x1": x1, "y1": y1})
        ln["words"].append(word)
        ln.update(x0=min(ln["x0"], x0), y0=min(ln["y0"], y0), x1=max(ln["x1"], x1), y1=max(ln["y1"], y1))
    out = []
    for ln in lines.values():
        ln["text"] = " ".join(ln["words"])
        ln["cx"], ln["cy"] = (ln["x0"] + ln["x1"]) / 2, (ln["y0"] + ln["y1"]) / 2
        out.append(ln)
    return out


def extract_labels(path):
    """Pair each equipment tag on a drawing with the caption written nearest to it.

    Returns {"pairs": [(tag, name)], "tags": [all tags]}. A vision model sees the shapes
    but often attaches the wrong tag to them; positions from OCR settle which is which.
    """
    with Image.open(path) as img:
        lines = _text_lines(img)
        width, height = _prepare(img).size
        block_text = _ocr_image(img, 6)
    tag_lines = [ln for ln in lines if _TAG.fullmatch(ln["text"])]
    captions = [ln for ln in lines
                if not _TAG.search(ln["text"]) and re.search(r"[A-Za-z]{4,}", ln["text"])
                and (ln["x1"] - ln["x0"]) < 0.4 * width]

    candidates = []
    for t in tag_lines:
        for c in captions:
            dx, dy = abs(c["cx"] - t["cx"]), c["cy"] - t["cy"]
            if dx < 0.08 * width and 0 < dy < 0.36 * height:  # caption sits below its tag
                candidates.append((dx * 3 + dy, id(t), id(c), t, c))
    pairs, used_tags, used_caps = [], set(), set()
    for _, tid, cid, t, c in sorted(candidates, key=lambda x: x[0]):
        if tid in used_tags or cid in used_caps:
            continue
        used_tags.add(tid)
        used_caps.add(cid)
        name = c["text"]
        for extra in captions:  # a second caption line directly underneath ("Distillation" / "Column")
            if (id(extra) not in used_caps and 0 < extra["y0"] - c["y1"] < 1.2 * (c["y1"] - c["y0"])
                    and abs(extra["cx"] - c["cx"]) < 0.05 * width):
                name += " " + extra["text"]
                used_caps.add(id(extra))
        pairs.append((t["text"], name))
    tags = sorted({ln["text"] for ln in tag_lines} | set(_TAG.findall(block_text)))
    return {"pairs": sorted(pairs), "tags": tags}


def extract_text(path):
    """Returns {"text", "method", "pages"} for a PDF, image, DOCX or text file."""
    path = Path(path)
    ext = path.suffix.lower()

    if ext in IMAGE_EXTS:
        with Image.open(path) as img:
            return {"text": _ocr_image(img).strip(), "method": "OCR (Tesseract)", "pages": 1}

    if ext == ".pdf":
        import pymupdf
        parts, used_ocr = [], False
        with pymupdf.open(path) as doc:
            pages = len(doc)
            for page in doc:
                text = page.get_text().strip()
                if len(text) < 40:  # no text layer: the page is a scan
                    pix = page.get_pixmap(dpi=200)
                    text = _ocr_image(Image.open(io.BytesIO(pix.tobytes("png")))).strip()
                    used_ocr = True
                parts.append(text)
        method = "OCR (Tesseract)" if used_ocr else "PDF text layer"
        return {"text": "\n\n".join(parts).strip(), "method": method, "pages": pages}

    if ext == ".docx":
        import docx
        d = docx.Document(str(path))
        lines = [p.text for p in d.paragraphs if p.text.strip()]
        for table in d.tables:
            for row in table.rows:
                lines.append(" | ".join(c.text.strip() for c in row.cells))
        return {"text": "\n".join(lines), "method": "Word parser", "pages": 1}

    return {"text": path.read_text(encoding="utf-8", errors="replace"), "method": "Plain text", "pages": 1}


def read_document(args, ctx):
    path = resolve_input(args.get("path", ""))
    result = extract_text(path)
    if not result["text"]:
        return {"ok": False, "output": f"No readable text found in {path.name}."}
    return {"ok": True,
            "output": f"[{path.name} | {result['method']} | {result['pages']} page(s)]\n{result['text']}",
            "method": result["method"]}
